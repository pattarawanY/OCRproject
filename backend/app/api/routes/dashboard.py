"""แดชบอร์ดภาพรวมและรายงาน (ผู้ดูแลระบบ)"""

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy import extract, func, select
from sqlalchemy.orm import Session

from app.core.deps import require_roles
from app.db import get_db
from app.models import RequestStatus as S, Role, TeacherProfile, TravelRequest, User
from app.services.pdf import STATUS_TH, build_report_pdf

router = APIRouter(prefix="/admin/dashboard", tags=["admin: dashboard"])
admin_only = require_roles(Role.ADMIN)


def fiscal_year_range(fy: int) -> tuple[date, date]:
    """ปีงบประมาณ (พ.ศ.) เริ่ม 1 ต.ค. ของปีก่อนหน้า"""
    ce = fy - 543
    return date(ce - 1, 10, 1), date(ce, 9, 30)


def compute_stats(db: Session, fiscal_year: int | None, month: int | None, status: str | None) -> dict:
    filters = [TravelRequest.status != S.DRAFT]
    if fiscal_year:
        start, end = fiscal_year_range(fiscal_year)
        filters += [TravelRequest.start_date >= start, TravelRequest.start_date <= end]
    if month:
        filters.append(extract("month", TravelRequest.start_date) == month)
    if status:
        filters.append(TravelRequest.status == status)

    by_status = dict.fromkeys(S.ALL[1:], 0)
    for st, cnt in db.execute(
        select(TravelRequest.status, func.count()).where(*filters).group_by(TravelRequest.status)
    ):
        by_status[st] = cnt

    total_reimbursed = db.scalar(
        select(func.coalesce(func.sum(TravelRequest.actual_amount), 0)).where(
            *filters, TravelRequest.status == S.COMPLETED
        )
    ) or Decimal(0)
    total_estimated = db.scalar(
        select(func.coalesce(func.sum(TravelRequest.estimated_amount), 0)).where(*filters)
    ) or Decimal(0)

    budgets = [
        {
            "user_id": u.user_id,
            "full_name": u.full_name,
            "budget_total": float(p.budget_total),
            "budget_used": float(p.budget_used),
            "budget_remaining": float(p.budget_remaining),
        }
        for u, p in db.execute(
            select(User, TeacherProfile).join(TeacherProfile).where(User.role == Role.TEACHER).order_by(User.full_name)
        )
    ]
    users_by_role = {r: 0 for r in Role.ALL}
    for role, cnt in db.execute(select(User.role, func.count()).where(User.is_active.is_(True)).group_by(User.role)):
        users_by_role[role] = cnt

    return {
        "by_status": by_status,
        "request_count": sum(by_status.values()),
        "total_reimbursed": float(total_reimbursed),
        "total_estimated": float(total_estimated),
        "budgets": budgets,
        "user_count": sum(users_by_role.values()),
        "users_by_role": users_by_role,
    }


@router.get("")
def dashboard(
    fiscal_year: int | None = Query(None, description="ปีงบประมาณ พ.ศ. เช่น 2569"),
    month: int | None = Query(None, ge=1, le=12),
    status: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(admin_only),
):
    return compute_stats(db, fiscal_year, month, status)


@router.get("/report")
def download_report(
    format: Literal["pdf", "xlsx"] = "pdf",
    fiscal_year: int | None = None,
    month: int | None = Query(None, ge=1, le=12),
    status: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(admin_only),
):
    stats = compute_stats(db, fiscal_year, month, status)
    title = "รายงานภาพรวมการเบิกจ่าย CS SmartClaim" + (f" ปีงบประมาณ {fiscal_year}" if fiscal_year else "")
    stamp = datetime.now().strftime("%Y%m%d")

    if format == "pdf":
        return Response(
            build_report_pdf(stats, title),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="smartclaim_report_{stamp}.pdf"'},
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    ws.append([title])
    ws.append(["จำนวนผู้ใช้งาน", stats["user_count"]])
    ws.append(["จำนวนคำขอ", stats["request_count"]])
    ws.append(["ยอดประมาณการรวม (บาท)", stats["total_estimated"]])
    ws.append(["ยอดเบิกจ่ายรวม (บาท)", stats["total_reimbursed"]])
    ws.append([])
    ws.append(["สถานะ", "จำนวน"])
    for k, v in stats["by_status"].items():
        ws.append([STATUS_TH.get(k, k), v])

    ws2 = wb.create_sheet("วงเงินอาจารย์")
    ws2.append(["อาจารย์", "วงเงินทั้งหมด", "ใช้ไป", "คงเหลือ"])
    for b in stats["budgets"]:
        ws2.append([b["full_name"], b["budget_total"], b["budget_used"], b["budget_remaining"]])

    ws3 = wb.create_sheet("รายการคำขอ")
    ws3.append(["เลขที่", "อาจารย์", "กิจกรรม", "ชื่องาน", "สถานที่", "เริ่ม", "สิ้นสุด", "สถานะ", "ประมาณการ", "ยอดจริง"])
    stmt = select(TravelRequest).where(TravelRequest.status != S.DRAFT).order_by(TravelRequest.request_id)
    if status:
        stmt = stmt.where(TravelRequest.status == status)
    for r in db.scalars(stmt):
        ws3.append([r.request_id, r.user.full_name, r.activity_type, r.event_name, r.location,
                    r.start_date.isoformat(), r.end_date.isoformat(), STATUS_TH.get(r.status, r.status),
                    float(r.estimated_amount or 0), float(r.actual_amount or 0)])

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="smartclaim_report_{stamp}.xlsx"'},
    )
